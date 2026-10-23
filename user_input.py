import openpyxl

wb = openpyxl.Workbook() #created brand new excel workbook
ws = wb.active #ws = worksheet, new excel wb comes with one ws
ws.title = "Holdings"

ws["A1"] = "Asset Name"
ws["B1"] = "Ticker"
ws["C1"] = "Amount Invested (£)" 

example = [
    ("Global equity ETF", "VWRL.L", 5000),
    ("UK gilts ETF",      "IGLT.L", 3000),
    ("US tech ETF",       "EQQQ.L", 2000),
]

for i, (name, ticker, amount) in enumerate(example, start = 2):
    ws[f"A{i}"] = name
    ws[f"B{i}"] = ticker
    ws[f"C{i}"] = amount 

wb.save("holdings.xlsx")
print("Created holdings.xlsx")